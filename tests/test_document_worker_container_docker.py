"""F02 container profile, qualified against a real Docker engine and image.

Opt-in: set ``KAZMA_TEST_DOCUMENT_WORKER_IMAGE`` to a built worker image
(``docker build -f deploy/document-worker/Dockerfile -t <tag> .``). Without
it every test here is skipped -- a skip is not evidence. The tests drive the
production paths: the broker's fixed policy and attestation, its TCP
server, Kazma's client, ``DocumentService`` and ``DocumentOperations``.
Containers are labelled per broker name and the module checks none is left.
"""

from __future__ import annotations

import json
import os
import socket
import struct
import subprocess
import threading
import time
import uuid
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest
from kazma_core.docker_cli import find_docker_cli
from kazma_core.documents import worker_broker as broker_module
from kazma_core.documents import worker_container
from kazma_core.documents.config import DocumentConfig
from kazma_core.documents.errors import DocumentParseError
from kazma_core.documents.hostile_corpus import write_hostile_corpus
from kazma_core.documents.service import DocumentService
from kazma_core.documents.worker_broker import BrokerLimits, BrokerSettings, _Broker

IMAGE = (os.environ.get("KAZMA_TEST_DOCUMENT_WORKER_IMAGE") or "").strip()
DOCKER = find_docker_cli() if IMAGE else None
pytestmark = pytest.mark.skipif(
    not (IMAGE and DOCKER),
    reason="set KAZMA_TEST_DOCUMENT_WORKER_IMAGE (and have a docker CLI) to qualify the container profile",
)
# Errnos the probe reports from INSIDE the Linux container (the host's errno
# module differs on Windows: ENETUNREACH is 10051 there).
LINUX_EAGAIN, LINUX_EACCES, LINUX_ENOSPC, LINUX_EROFS, LINUX_ENETUNREACH = 11, 13, 28, 30, 101
PROD_NAME = "f02prod" + uuid.uuid4().hex[:6]
QUAL_NAME = "f02qual" + uuid.uuid4().hex[:6]


def _labelled(name: str) -> list[str]:
    result = subprocess.run(
        [DOCKER, "ps", "-a", "-q", "--filter", f"label={broker_module.LABEL}={name}"],
        capture_output=True, text=True, check=True, timeout=30,
    )
    return result.stdout.split()


def _start(broker: _Broker) -> tuple[Any, threading.Thread]:
    server = broker_module._make_server("tcp:127.0.0.1:0", broker, allow_non_loopback=False)
    thread = threading.Thread(target=server.serve_forever, kwargs={"poll_interval": 0.05}, daemon=True)
    thread.start()
    return server, thread


@pytest.fixture(scope="module")
def token_file(tmp_path_factory: pytest.TempPathFactory) -> Path:
    path = tmp_path_factory.mktemp("broker") / "token"
    path.write_text(uuid.uuid4().hex + uuid.uuid4().hex, encoding="utf-8")
    return path


@pytest.fixture(scope="module")
def production_broker(token_file: Path):
    """A broker with the shipped defaults, serving on loopback."""

    broker = _Broker(BrokerSettings(
        docker=DOCKER, image=IMAGE, token=token_file.read_text(encoding="utf-8"), name=PROD_NAME,
    ))
    broker.prepare()
    assert broker.attestation["verified"], broker.attestation["problems"]
    assert broker.capabilities is not None, broker.capability_error
    server, thread = _start(broker)
    yield broker, server
    server.shutdown()
    server.server_close()
    thread.join(timeout=10)
    broker.shutdown()
    assert _labelled(PROD_NAME) == []


@pytest.fixture(scope="module")
def qualification_broker(token_file: Path):
    """Small limits and the probe's stress actions, to watch each limit hold."""

    limits = BrokerLimits(memory_mb=256, cpus=1.0, pids=64, tmpfs_mb=64, concurrency=1,
                          max_timeout_seconds=60, max_output_bytes=4 * 1024 * 1024,
                          admission_wait_seconds=1.0)
    broker = _Broker(BrokerSettings(
        docker=DOCKER, image=IMAGE, token=token_file.read_text(encoding="utf-8"),
        name=QUAL_NAME, limits=limits, qualification=True,
    ))
    broker.prepare()
    assert broker.attestation["verified"], broker.attestation["problems"]
    assert broker.capabilities is not None, broker.capability_error
    server, thread = _start(broker)
    yield broker, server
    server.shutdown()
    server.server_close()
    thread.join(timeout=10)
    broker.shutdown()
    assert _labelled(QUAL_NAME) == []


@pytest.fixture
def container_profile(production_broker, token_file: Path, monkeypatch: pytest.MonkeyPatch):
    _, server = production_broker
    host, port = server.server_address[:2]
    monkeypatch.setenv("KAZMA_DOCUMENT_WORKER_ISOLATION", "container")
    monkeypatch.setenv("KAZMA_DOCUMENT_BROKER_URL", f"tcp://{host}:{port}")
    monkeypatch.setenv("KAZMA_DOCUMENT_BROKER_TOKEN_FILE", str(token_file))
    monkeypatch.setattr(worker_container, "_status_cache", None)
    return production_broker[0]


# ── the policy holds inside a real container ────────────────────────────


def test_attestation_from_inside_the_job_container(production_broker) -> None:
    broker, _ = production_broker
    facts = broker.attestation["facts"]
    assert facts["interfaces"] == ["lo"]
    assert facts["external_connect_errno"] == LINUX_ENETUNREACH
    assert facts["dns_resolved"] is False
    assert facts["uid"] == 10001 and facts["status"]["CapEff"] == "0000000000000000"
    assert facts["status"]["NoNewPrivs"] == "1" and facts["status"]["Seccomp"] == "2"
    assert set(facts["write_errno"].values()) <= {LINUX_EROFS, LINUX_EACCES}
    assert facts["sensitive_paths_visible"] == []
    points = {row["mount_point"] for row in facts["mounts"]}
    assert not any(p.startswith(("/app", "/home", "/kazma", "/var/run", "/run/docker")) for p in points)
    assert facts["cgroup"]["memory.swap.max"] == "0"
    assert broker.attestation["cleanup"] == "verified"


def test_the_isolation_report_reads_the_attested_broker(container_profile) -> None:
    report = worker_container.isolation_report()
    assert report["mode"] == "container" and report["ready"] is True
    assert report["filesystem_isolation"] and report["network_isolation"]
    assert report["broker"]["image_id"] == container_profile.image_id


# ── real documents through the production paths ─────────────────────────


def _fixtures(root: Path) -> dict[str, Path]:
    from docx import Document
    from openpyxl import Workbook
    from reportlab.pdfgen import canvas

    root.mkdir(parents=True, exist_ok=True)
    pdf = root / "english.pdf"
    page = canvas.Canvas(str(pdf))
    page.drawString(72, 720, "Kazma container parse check: invoice 2026-10-13 total 4,250.75")
    page.showPage()
    page.drawString(72, 720, "Second page keeps its own number.")
    page.save()
    docx = root / "bilingual.docx"
    document = Document()
    document.add_heading("Quarterly report", level=1)
    document.add_paragraph("English paragraph with a number 1,234.")
    document.add_paragraph("فقرة عربية تذكر الرقم ١٢٣٤ والتاريخ.")
    table = document.add_table(rows=2, cols=2)
    for r, row in enumerate([("Item", "البند"), ("Total", "٤٢٥٠")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.save(str(docx))
    xlsx = root / "sheet.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "الاسم", "amount"])
    sheet.append(["alpha", "ألفا", 10.5])
    book.save(str(xlsx))
    csv = root / "table.csv"
    csv.write_text("a,b\n1,٢\n", encoding="utf-8")
    md = root / "notes.md"
    md.write_text("# Notes\n\nPlain **markdown** with Arabic: مرحبا.\n", encoding="utf-8")
    return {"pdf": pdf, "docx": docx, "xlsx": xlsx, "csv": csv, "md": md}


def _texts(ir: Any) -> list[tuple[int, str]]:
    return [(page.page_number, block.text) for page in ir.pages for block in page.blocks]


def _config(root: Path, **overrides: Any) -> DocumentConfig:
    return replace(DocumentConfig(storage_root=root / "store"), ocr_enabled=False,
                   security_local_salvage=False, **overrides)


async def _parse(service: DocumentService, path: Path) -> Any:
    return await service.parse_path(path, tenant_id="t", workspace_id="w", actor_id="a")


async def test_container_parses_match_the_local_worker(container_profile, tmp_path: Path, monkeypatch) -> None:
    fixtures = _fixtures(tmp_path / "docs")
    service = DocumentService(config=_config(tmp_path))
    compared = []
    for kind, path in fixtures.items():
        contained = await _parse(service, path)
        monkeypatch.setenv("KAZMA_DOCUMENT_WORKER_ISOLATION", "process")
        local = await _parse(service, path)
        monkeypatch.setenv("KAZMA_DOCUMENT_WORKER_ISOLATION", "container")
        assert _texts(contained) == _texts(local), kind
        assert contained.metadata.get("source_sha256") == local.metadata.get("source_sha256")
        compared.append((kind, len(_texts(contained))))
    assert all(count > 0 for _, count in compared), compared
    joined = " ".join(text for _, text in _texts(await _parse(service, fixtures["docx"])))
    assert "فقرة عربية" in joined and "1,234" in joined


def _outcome(call) -> str:
    try:
        call()
    except DocumentParseError as exc:
        return exc.code
    return "ok"


async def test_hostile_corpus_outcomes_match_in_the_container(container_profile, tmp_path: Path, monkeypatch) -> None:
    manifest_path, manifest = write_hostile_corpus(tmp_path / "corpus")
    root = manifest_path.parent
    results = []
    for case in manifest["cases"]:
        if case["stage"] == "sandbox":
            continue  # injected parser crash: a monkeypatch of the local runner
        service = DocumentService(config=_config(tmp_path / case["id"], **case["config_overrides"]))
        path = root / case["filename"]

        def run(service=service, path=path, stage=case["stage"]):
            if stage == "sniff":
                service.inspect_path(path)
            else:
                service._parse_isolated(path)

        contained = _outcome(run)
        monkeypatch.setenv("KAZMA_DOCUMENT_WORKER_ISOLATION", "process")
        local = _outcome(run)
        monkeypatch.setenv("KAZMA_DOCUMENT_WORKER_ISOLATION", "container")
        expected = set(case["expected_codes"]) or {"ok"}
        results.append((case["id"], contained, local))
        assert contained == local, (case["id"], contained, local)
        assert contained in expected, (case["id"], contained, expected)
    assert len(results) >= 18, results


async def test_mutation_and_render_workers_run_in_the_container(container_profile, tmp_path: Path) -> None:
    from pypdf import PdfReader, PdfWriter

    sources = []
    for index in range(2):
        path = tmp_path / f"{index}.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        with path.open("wb") as handle:
            writer.write(handle)
        sources.append(path)
    service = DocumentService(config=_config(tmp_path))
    merged = await service.pdf_merge(
        tuple(sources), approved_paths=tuple(sources), tenant_id="t", workspace_id="w",
        actor_id="a", export_dir=tmp_path / "exports",
    )
    assert merged.ok, (merged.code, merged.message)
    assert len(PdfReader(str(merged.data.export_path)).pages) == 2
    info = await service.pdf_info(sources[0], approved_path=sources[0], tenant_id="t", workspace_id="w", actor_id="a")
    assert info.ok and info.data["page_count"] == 1
    for target in ("docx", "markdown"):
        generated = await service.generate(
            target, {"title": "Report تقرير", "sections": [{"heading": "One", "body": "مرحبا Hello"}]},
            tenant_id="t", workspace_id="w", actor_id="a", export_dir=tmp_path / "exports",
        )
        assert generated.ok, (target, generated.code, generated.message)
        assert generated.data.export_path.stat().st_size > 0


# ── each limit holds (qualification broker) ─────────────────────────────


def _run(broker: _Broker, request: dict[str, Any], timeout: float = 60.0) -> tuple[dict, dict | None]:
    status, result = broker.run_local("probe", request, timeout_seconds=timeout)
    assert status["cleanup"] == "verified", status
    return status, result


def test_memory_limit_kills_and_reports(qualification_broker) -> None:
    status, _ = _run(qualification_broker[0], {"action": "allocate", "mib": 512})
    assert status["status"] == "memory_limit" and status["oom_killed"] is True


def test_process_count_is_capped(qualification_broker) -> None:
    status, result = _run(qualification_broker[0], {"action": "spawn", "count": 500})
    assert status["status"] == "completed", status
    facts = result["facts"]
    assert facts["spawned"] < 64 and facts["errno"] == LINUX_EAGAIN


def test_cpu_is_capped_at_one_core(qualification_broker) -> None:
    status, result = _run(qualification_broker[0], {"action": "spin", "workers": 4, "seconds": 4})
    assert status["status"] == "completed", status
    assert result["facts"]["cores_used"] <= 1.15, result["facts"]


def test_scratch_space_is_bounded(qualification_broker) -> None:
    status, result = _run(qualification_broker[0], {"action": "fill_tmp"})
    assert status["status"] == "completed", status
    facts = result["facts"]
    assert facts["errno"] == LINUX_ENOSPC and facts["written"] <= 64 * 1024 * 1024


def test_deadline_kills_and_removes(qualification_broker) -> None:
    started = time.monotonic()
    status, _ = _run(qualification_broker[0], {"action": "sleep", "seconds": 120}, timeout=3)
    assert status["status"] == "timeout"
    assert time.monotonic() - started < 30


def test_output_flood_is_cut_at_the_cap(qualification_broker) -> None:
    status, _ = _run(qualification_broker[0], {"action": "flood", "mib": 16})
    assert status["status"] == "output_limit"
    assert status["output_bytes"] > 4 * 1024 * 1024


def test_stderr_is_retained_only_to_its_bound(qualification_broker) -> None:
    status, _ = _run(qualification_broker[0], {"action": "stderr_flood", "mib": 4})
    assert status["status"] == "completed"
    assert status["stderr_truncated"] is True
    assert len(status["stderr"].encode("utf-8")) <= 64 * 1024


def test_a_detached_descendant_dies_with_its_container(qualification_broker) -> None:
    status, result = _run(qualification_broker[0], {"action": "orphan"})
    assert status["status"] == "completed" and result["facts"]["orphan_pid"] > 0
    assert _labelled(QUAL_NAME) == []


def test_a_client_that_disconnects_cancels_the_job(qualification_broker, token_file: Path) -> None:
    broker, server = qualification_broker
    payload = broker_module._bundle_bytes(
        broker_module._INPUT_MAGIC,
        [("work/request.json", json.dumps({"action": "sleep", "seconds": 120}).encode())],
        {"worker": "probe"},
    )
    run_id = uuid.uuid4().hex
    header = json.dumps({"protocol": 1, "token": token_file.read_text(encoding="utf-8"), "op": "run",
                         "worker": "probe", "run_id": run_id, "timeout_seconds": 60,
                         "input_size": len(payload)}).encode()
    sock = socket.create_connection(server.server_address[:2], timeout=30)
    sock.sendall(broker_module.REQUEST_MAGIC + struct.pack(">I", len(header)) + header)
    assert sock.recv(5)[:1] == b"A"
    sock.sendall(payload)
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline and not _container_for(run_id):
        time.sleep(0.2)
    assert _container_for(run_id), "job container never started"
    started = time.monotonic()
    sock.close()
    while time.monotonic() - started < 30 and _container_for(run_id):
        time.sleep(0.2)
    assert not _container_for(run_id)
    assert time.monotonic() - started < 30


def _container_for(run_id: str) -> list[str]:
    result = subprocess.run(
        [DOCKER, "ps", "-a", "-q", "--filter", f"label={broker_module.RUN_LABEL}={run_id}"],
        capture_output=True, text=True, check=True, timeout=30,
    )
    return result.stdout.split()


def test_a_full_broker_answers_busy(qualification_broker, token_file: Path) -> None:
    broker, _ = qualification_broker
    holder = threading.Thread(target=broker.run_local, args=("probe", {"action": "sleep", "seconds": 8}),
                              kwargs={"timeout_seconds": 30})
    holder.start()
    time.sleep(2.0)
    try:
        assert broker.acquire() is False
    finally:
        holder.join(timeout=60)


def test_startup_sweep_removes_containers_left_by_an_earlier_run(qualification_broker) -> None:
    broker, _ = qualification_broker
    stale = subprocess.run(
        [DOCKER, "run", "-d", "--network", "none", "--label", f"{broker_module.LABEL}={QUAL_NAME}",
         "--entrypoint", "sleep", broker.image_id, "300"],
        capture_output=True, text=True, check=True, timeout=60,
    ).stdout.strip()
    assert stale and _labelled(QUAL_NAME)
    assert broker.sweep() == 1
    assert _labelled(QUAL_NAME) == []


def test_a_production_broker_refuses_stress_actions(production_broker) -> None:
    status, result = production_broker[0].run_local("probe", {"action": "sleep", "seconds": 1})
    assert status["status"] == "completed"
    assert result["code"] == "probe_action_refused"


def test_a_worker_crash_reaches_kazma_as_no_result(qualification_broker, token_file: Path, tmp_path: Path, monkeypatch) -> None:
    _, server = qualification_broker
    host, port = server.server_address[:2]
    monkeypatch.setenv("KAZMA_DOCUMENT_BROKER_URL", f"tcp://{host}:{port}")
    monkeypatch.setenv("KAZMA_DOCUMENT_BROKER_TOKEN_FILE", str(token_file))
    result = worker_container.run_container_worker(
        worker="probe", request={"action": "crash"}, inputs=[], run_dir=tmp_path,
        timeout_seconds=30, max_output_bytes=1024 * 1024,
    )
    assert result.returncode == 139  # SIGSEGV, reported through --init
    assert not (tmp_path / "result.json").exists()
    assert _labelled(QUAL_NAME) == []


def test_a_missing_worker_image_refuses_to_start(token_file: Path) -> None:
    broker = _Broker(BrokerSettings(
        docker=DOCKER, image="kazma-document-worker:missing-" + uuid.uuid4().hex[:8],
        token=token_file.read_text(encoding="utf-8"), name="f02missing",
    ))
    with pytest.raises(RuntimeError, match="not available locally"):
        broker.resolve_image()
    assert broker.image_id is None and broker.status()["ready"] is False


def test_the_image_reports_its_own_engines(container_profile) -> None:
    from kazma_core.documents.worker_bundle import code_fingerprint

    report = container_profile.capabilities
    assert report is not None, container_profile.capability_error
    parsers = {row["parser_id"]: row["readiness"] for row in report["parsers"]}
    assert parsers["pdf"] == "ready" and parsers["docx"] == "ready"
    assert parsers["legacy-doc"] == "ready"  # LibreOffice is in the image
    renderers = {row["renderer_id"]: row["readiness"] for row in report["renderers"]}
    assert renderers["weasyprint"] == "ready" and renderers["libreoffice"] == "ready"
    assert {"ara", "eng"} <= set(report["ocr"]["languages"])
    local = Path(worker_container.__file__).resolve().parent
    assert report["documents_code_sha256"] == code_fingerprint(local)
    assert worker_container.isolation_report()["image_code_matches"] is True


def test_the_neural_ocr_engine_reads_inside_the_container(container_profile, tmp_path: Path) -> None:
    """2026-10-15 (Q02): the image carries the neural OCR engine and its
    models (no network at run time) and reads an Arabic-Indic number with it.
    Negative control: the same page with documents.ocr.engine: tesseract."""
    # The page is drawn here, in the test process: PyMuPDF is the documents
    # extra's, not the test extra's (the CI job installs it; a skip there
    # fails the job).
    pymupdf = pytest.importorskip("pymupdf")

    report = container_profile.capabilities
    assert report["ocr"]["engines"]["neural"]["readiness"] == "ready", report["ocr"]["engines"]
    fonts = Path(worker_container.__file__).resolve().parent / "assets" / "fonts"
    text = "بلغت الإيرادات لهذا الفصل ١٢٣٤٥٦٧ ريالا في جميع المناطق"
    css = "@font-face {font-family: F; src: url(IBMPlexSansArabic-Regular.ttf);} body {font-family: F; font-size: 14pt;}"
    story = pymupdf.Story(html=f'<body dir="rtl"><p dir="rtl">{text}</p></body>', user_css=css,
                          archive=pymupdf.Archive(str(fonts)))
    pdf = tmp_path / "line.pdf"
    writer = pymupdf.DocumentWriter(str(pdf))
    device = writer.begin_page(pymupdf.Rect(0, 0, 595, 120))
    story.place(pymupdf.Rect(30, 30, 565, 100))
    story.draw(device)
    writer.end_page()
    writer.close()
    image = tmp_path / "line.png"
    with pymupdf.open(pdf) as doc:
        doc[0].get_pixmap(dpi=300).save(image)

    def read(engine: str):
        service = DocumentService(config=replace(_config(tmp_path / engine), ocr_enabled=True, ocr_engine=engine))
        return service.parse_ingested_blob(image, force_ocr=True, ocr_language="ara")

    neural = read("neural")
    assert neural.metadata["ocr"]["engine"] == "neural"
    assert "١٢٣٤٥٦٧" in " ".join(text for _, text in _texts(neural))
    tesseract = read("tesseract")
    assert tesseract.metadata["ocr"]["engine"] == "tesseract"
    assert "١٢٣٤٥٦٧" not in " ".join(text for _, text in _texts(tesseract))


async def test_readiness_follows_the_image_not_this_host(container_profile, tmp_path: Path, monkeypatch) -> None:
    from kazma_core.documents.ocr import get_ocr_health
    from kazma_core.documents.renderers import get_renderer_registry

    assert get_renderer_registry().resolve("convert:html:pdf").available
    assert get_ocr_health(("ara", "eng")).readiness.value in {"ready", "degraded"}
    html = tmp_path / "page.html"
    html.write_text(
        "<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>Report</title></head>"
        "<body><h1>Report تقرير</h1><p>Body text مرحبا.</p></body></html>",
        encoding="utf-8",
    )
    service = DocumentService(config=_config(tmp_path))
    converted = await service.convert(
        html, approved_path=html, target_format="pdf", tenant_id="t", workspace_id="w",
        actor_id="a", export_dir=tmp_path / "exports",
    )
    assert converted.ok, (converted.code, converted.message)
    assert converted.data.export_path.read_bytes().startswith(b"%PDF")
    # This host's own view, recorded rather than asserted (it varies by host).
    monkeypatch.setenv("KAZMA_DOCUMENT_WORKER_ISOLATION", "process")
    print("host weasyprint:", get_renderer_registry().resolve("convert:html:pdf").readiness.value)


async def test_weasyprint_fetches_approved_assets_and_refuses_the_rest(container_profile, tmp_path: Path) -> None:
    from PIL import Image

    logo = tmp_path / "logo.png"
    Image.new("RGB", (8, 8), (200, 30, 30)).save(logo)
    service = DocumentService(config=_config(tmp_path))
    page = tmp_path / "with-asset.html"
    page.write_text(
        "<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>A</title></head>"
        "<body><p>Logo:</p><img src=\"logo.png\" alt=\"logo\"></body></html>",
        encoding="utf-8",
    )
    allowed = await service.convert(
        page, approved_path=page, target_format="pdf", tenant_id="t", workspace_id="w",
        actor_id="a", export_dir=tmp_path / "exports", approved_assets=(logo,),
    )
    assert allowed.ok, (allowed.code, allowed.message)
    external = tmp_path / "external.html"
    external.write_text(
        "<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>E</title></head>"
        "<body><img src=\"https://example.com/tracker.png\" alt=\"x\"></body></html>",
        encoding="utf-8",
    )
    refused = await service.convert(
        external, approved_path=external, target_format="pdf", tenant_id="t", workspace_id="w",
        actor_id="a", export_dir=tmp_path / "exports",
    )
    assert not refused.ok and refused.code == "external_resource_denied", (refused.code, refused.message)
