"""Behavioral recovery/fidelity regressions from the October Documents audit."""

import io
import json
import os
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from kazma_core.documents.backup import perform_document_backup, verify_document_backup
from kazma_core.documents.config import DocumentConfig
from kazma_core.documents.errors import DocumentSandboxError, DocumentSecurityError
from kazma_core.documents.ingestion import DocumentIngestionService
from kazma_core.documents.models import DocumentJobState
from kazma_core.documents.service import DocumentService
from kazma_core.documents.worker import DocumentWorker


def _service(tmp_path):
    return DocumentIngestionService(config=DocumentConfig(
        storage_root=tmp_path / 'store', security_malware_scan='off',
        security_local_salvage=False))


def _ingest(svc):
    return svc.ingest_stream(io.BytesIO(b'An audit recovery fixture.'), filename='note.txt',
        tenant_id='tenant', workspace_id='workspace', actor_id='alice')


def test_backup_checks_metadata_and_normalized_manifest(tmp_path):
    svc = _service(tmp_path)
    try:
        result = _ingest(svc)
        svc.storage.write_manifest(document_id=result.document_id, version_id=result.version_id,
            manifest={'ir': {'pages': []}})
        report = perform_document_backup(storage_root=svc.config.storage_root, dest_dir=tmp_path / 'backups')
        assert report['ok']
        normalized = f'manifests/{result.document_id}/{result.version_id}.json'
        assert normalized in report['checksums'] and 'documents.db' in report['checksums']
        path = Path(report['path']) / normalized
        path.write_text('corrupt', encoding='utf-8')
        assert normalized in verify_document_backup(backup_dir=report['path'])['corrupt']
        path.unlink()
        assert normalized in verify_document_backup(backup_dir=report['path'])['missing']
    finally:
        svc.close()


def test_failed_backup_does_not_displace_verified_retention(tmp_path):
    from kazma_core.documents.backup import _prune_old
    svc = _service(tmp_path)
    try:
        _ingest(svc)
        reports = [perform_document_backup(storage_root=svc.config.storage_root,
            dest_dir=tmp_path / 'backups', retention=10) for _ in range(3)]
        paths = [Path(report['path']) for report in reports]
        for index, path in enumerate(paths):
            os.utime(path, (100 + index, 100 + index))
        # A failed directory newer than two good recovery points used to
        # consume a retention slot and evict the older good snapshot.
        (paths[2] / 'documents.db').write_bytes(b'corrupt snapshot')
        assert _prune_old(tmp_path / 'backups', retention=2) == 0
        assert all(path.exists() for path in paths)
        assert _prune_old(tmp_path / 'backups', retention=1) == 1
        assert not paths[0].exists() and paths[1].exists()
        assert verify_document_backup(backup_dir=paths[1])['ok']
    finally:
        svc.close()


@pytest.mark.parametrize('body', [None, '{', '{}', '{"checksums":{}}', '{"checksums":[]}'])
def test_backup_missing_or_invalid_inventory_fails_closed(tmp_path, body):
    if body is not None:
        (tmp_path / 'manifest.json').write_text(body, encoding='utf-8')
    assert not verify_document_backup(backup_dir=tmp_path)['ok']


def test_verifier_rejects_omitted_referenced_blob(tmp_path):
    svc = _service(tmp_path)
    try:
        _ingest(svc)
        report = perform_document_backup(storage_root=svc.config.storage_root, dest_dir=tmp_path / 'backups')
        assert report['ok']
        inventory = Path(report['path']) / 'manifest.json'
        value = json.loads(inventory.read_text(encoding='utf-8'))
        value['checksums'] = {'documents.db': value['checksums']['documents.db']}
        inventory.write_text(json.dumps(value), encoding='utf-8')
        check = verify_document_backup(backup_dir=report['path'])
        assert not check['ok'] and any('required_object' in item for item in check['invalid_inventory'])
    finally:
        svc.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('permanent', [False, True])
async def test_real_queue_preserves_parser_failure_classification(tmp_path, monkeypatch, permanent):
    svc = _service(tmp_path)
    now = datetime.now(UTC)
    monkeypatch.setattr(svc.jobs, '_clock', lambda: now)
    try:
        result = _ingest(svc)
        def fail(*args, **kwargs):
            if permanent:
                raise DocumentSecurityError('Unsafe content')
            raise DocumentSandboxError('Parser timed out', code='parser_timeout')
        monkeypatch.setattr(svc.service, 'parse_ingested_blob', fail)
        worker = DocumentWorker(svc.jobs, svc.stage_handlers())
        for attempt in range(1, 4):
            claimed = svc.jobs.claim_next(owner='audit-worker', lease_seconds=120)
            assert claimed is not None
            await worker._process_claimed(claimed, 'audit-worker')
            job = svc.jobs.get(tenant_id='tenant', job_id=result.job_id)
            if permanent:
                assert job.state is DocumentJobState.REJECTED
                break
            assert job.error_code == 'parser_timeout'
            assert job.state is (DocumentJobState.DEAD_LETTER if attempt == 3 else DocumentJobState.RETRY_WAIT)
            now += timedelta(hours=1)
    finally:
        svc.close()


@pytest.mark.asyncio
async def test_transient_intake_validation_is_reclaimable(tmp_path, monkeypatch):
    svc = _service(tmp_path)
    now = datetime.now(UTC)
    monkeypatch.setattr(svc.jobs, '_clock', lambda: now)
    validate = svc._validate_and_promote
    calls = 0
    def flaky(**kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise DocumentSandboxError('Preflight timed out', code='parser_timeout')
        return validate(**kwargs)
    monkeypatch.setattr(svc, '_validate_and_promote', flaky)
    try:
        result = _ingest(svc)
        assert result.state is DocumentJobState.RETRY_WAIT
        now += timedelta(hours=1)
        claimed = svc.jobs.claim_next(owner='audit-worker', lease_seconds=120)
        assert claimed.state is DocumentJobState.VALIDATING
        await DocumentWorker(svc.jobs, svc.stage_handlers())._process_claimed(claimed, 'audit-worker')
        assert svc.jobs.get(tenant_id='tenant', job_id=result.job_id).state is DocumentJobState.READY_TO_PARSE
        assert calls == 2
    finally:
        svc.close()


@pytest.mark.parametrize('before,inside,after', [
    ('Before the table.', 'Inside the table.', 'After the table.'),
    ('قبل الجدول', 'بيانات الجدول', 'بعد الجدول'),
])
def test_docx_body_order_survives_real_isolated_parse(tmp_path, before, inside, after):
    from docx import Document
    document = Document()
    document.add_paragraph(before)
    document.add_table(rows=1, cols=1).cell(0, 0).text = inside
    document.add_paragraph(after)
    path = tmp_path / 'ordered.docx'
    document.save(path)
    # Original boundary extracted document.paragraphs then document.tables.
    original = [p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert original == [before, after, inside]
    svc = DocumentService(config=DocumentConfig(storage_root=tmp_path / 'store',
        security_local_salvage=False, ocr_enabled=False))
    parsed = svc.read_transient_sync(path, approved_path=path, fence=False)
    assert parsed.text.index(before) < parsed.text.index(inside) < parsed.text.index(after)
