"""A generated Arabic PDF in the brand face, with its lines and levels kept.

Live 2026-10-18: the owner asked for a PDF of every X post. The file came
out of the document worker container in DejaVu Sans (Arabic) and Noto Serif
(Latin) -- the DOCX names IBM Plex Sans Arabic and LibreOffice in the image
had no Plex --, each tweet's lines ran into one justified paragraph, the
quotes were italic (Arabic slanted by LibreOffice), ``##`` groups and
``###`` entries were the same size, the heading bar of a Latin heading sat
on the right, the divider was a row of "─" in a fallback font, table rows
were tall with their text at the top, a date after an Arabic word drew as
"18-10-2026", and a deleted tweet's heading ("10 · 2026-10-05 — 15:15 UTC ·
<id> — محذوفة") came out in reverse group order.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from kazma_core.documents.content_model import BodyBlock, ContentModel, TitleBlock
from kazma_core.documents.engines.docx import DocxEngine
from kazma_core.documents.fonts import give_soffice_the_bundled_fonts
from kazma_core.documents.profile import DocProfile
from kazma_core.documents.rich_render import parse_rich_blocks

#: The live archive's shape (two tweets, a table, a divider, two levels);
#: ``<BR>`` is a hard break (two trailing spaces an editor would strip).
ARCHIVE = """# أرشيف تغريدات كاظمه على X

الحساب: @KazmaAI · المصدر: موصّل X الرسمي

---

## نطاق الملف — اقرأه أولاً

| البند | القيمة |
|---|---|
| منها منشورة الآن | 9 |

## أولاً: المنشورات القائمة (9)

### 1 · 2026-10-16 — 12:25 UTC · 2111071024812847259

https://x.com/KazmaAI/status/2111071024812847259

> Kazma metrics — 2026-10-16, commit 03f80393<BR>
> Python files: 2,371<BR>
> Collected test cases: 17,128 (collected, not passes)

### 2 · 2026-10-16 — 11:48 UTC · 2111061637000011943

> كاظمه بالأرقام 2026-10-16<BR>
> ملفات بايثون: 2,368<BR>
> #كاظمه
""".replace("<BR>", "  ")  # two trailing spaces: a Markdown hard break


# ── the Markdown: hard line breaks are kept ─────────────────────────────


def test_a_tweets_lines_stay_lines():
    quotes = [b["text"] for b in parse_rich_blocks(ARCHIVE) if b["type"] == "quote"]
    assert quotes[0].split("\n") == [
        "Kazma metrics — 2026-10-16, commit 03f80393",
        "Python files: 2,371",
        "Collected test cases: 17,128 (collected, not passes)",
    ]
    assert quotes[1].split("\n")[-1] == "#كاظمه"


def test_markdown_breaks_and_joins():
    assert parse_rich_blocks("one\ntwo")[0]["text"] == "one two"  # a soft newline is a space
    assert parse_rich_blocks("one  \ntwo\\\nthree")[0]["text"] == "one\ntwo\nthree"
    assert parse_rich_blocks("> x\n>\n> y")[0]["text"] == "x\ny"
    assert parse_rich_blocks("> a\n> b")[0]["text"] == "a b"


def test_negative_control_the_old_join_lost_the_lines():
    lines = [ln.lstrip("> ").strip() for ln in ARCHIVE.splitlines() if ln.startswith(">")][:3]
    assert "\n" not in " ".join(lines)


# ── the payload: a heading's level is its own ───────────────────────────


def _model(sections, title="أرشيف تغريدات كاظمه على X"):
    from kazma_core.documents.content_model import HeadingBlock
    from kazma_core.documents.renderer_worker import _build_model_and_profile

    model, _ = _build_model_and_profile({"title": title, "sections": sections, "lang": "ar", "rtl": True})
    return [(b.text, b.level) for b in model.blocks if isinstance(b, HeadingBlock)]


def test_section_levels_follow_their_hashes():
    sections = [{"heading": "# أرشيف تغريدات كاظمه على X", "body": "x"},
                {"heading": "## أولاً: المنشورات", "body": "x"},
                {"heading": "### 1 · 2026-10-16", "body": "x"},
                {"heading": "## ثانياً", "body": "x"},
                {"heading": "Plain heading", "body": "x"}]
    assert _model(sections) == [("أولاً: المنشورات", 1), ("1 · 2026-10-16", 2), ("ثانياً", 1),
                                ("Plain heading", 1)]


# ── the DOCX ────────────────────────────────────────────────────────────


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx(tmp_path: Path, text: str = ARCHIVE):
    from lxml import etree

    model = ContentModel()
    model.add(TitleBlock(text="أرشيف تغريدات كاظمه على X", level=0))
    model.add(BodyBlock(text=text))
    out = tmp_path / "a.docx"
    DocxEngine(DocProfile.for_content(text)).render(model, out)
    with zipfile.ZipFile(out) as z:
        return etree.fromstring(z.read("word/document.xml"))


def _paragraphs(root):
    for p in root.iter(f"{W}p"):
        text = "".join(t.text or "" for t in p.iter(f"{W}t"))
        yield p, text


def _side(p) -> str | None:
    bdr = p.find(f"{W}pPr/{W}pBdr")
    if bdr is None:
        return None
    sides = [c.tag.replace(W, "") for c in bdr if c.tag in (f"{W}left", f"{W}right")]
    return sides[0] if sides else None


def test_quotes_are_upright_broken_into_lines_and_barred_on_their_side(tmp_path):
    root = _docx(tmp_path)
    quotes = {text: p for p, text in _paragraphs(root) if text.startswith(("Kazma metrics", "كاظمه بالأرقام"))}
    en, ar = quotes["Kazma metrics — 2026-10-16, commit 03f80393Python files: 2,371Collected test cases: "
                    "17,128 (collected, not passes)"], next(p for t, p in quotes.items() if t.startswith("كاظمه"))
    for p in (en, ar):
        assert not list(p.iter(f"{W}i")), "no italic"
        assert len(list(p.iter(f"{W}br"))) == 2, "the tweet's line breaks"
    assert _side(en) == "left" and _side(ar) == "right"


def test_a_latin_headings_bar_sits_on_its_left_in_an_arabic_document(tmp_path):
    root = _docx(tmp_path, ARCHIVE.replace("### 1 · 2026-10-16", "## 1 · 2026-10-16"))
    latin = next(p for p, t in _paragraphs(root) if t.startswith("1 · 2026-10-16"))
    arabic = next(p for p, t in _paragraphs(root) if t.startswith("أولاً"))
    assert _side(latin) == "left" and _side(arabic) == "right"


def test_the_divider_is_a_rule_not_a_line_of_glyphs(tmp_path):
    root = _docx(tmp_path)
    assert not any("─" in t for _, t in _paragraphs(root))
    rules = [p for p, t in _paragraphs(root) if not t and p.find(f"{W}pPr/{W}pBdr/{W}bottom") is not None]
    assert rules


def test_borders_sit_where_the_schema_puts_them(tmp_path):
    """Word reports a file whose pPr children are out of order as damaged."""
    root = _docx(tmp_path)
    for p, _ in _paragraphs(root):
        ppr = p.find(f"{W}pPr")
        if ppr is None or ppr.find(f"{W}pBdr") is None:
            continue
        tags = [c.tag.replace(W, "") for c in ppr]
        before = tags[:tags.index("pBdr")]
        later = {t.replace("w:", "") for t in DocxEngine._AFTER_PBDR}
        assert not (later & set(before)), tags
        sides = [c.tag.replace(W, "") for c in ppr.find(f"{W}pBdr")]
        order = ["top", "left", "bottom", "right", "between", "bar"]
        assert sides == sorted(sides, key=order.index), sides


def test_table_cells_are_compact_and_centred(tmp_path):
    root = _docx(tmp_path)
    tc = next(root.iter(f"{W}tc"))
    assert tc.find(f"{W}tcPr/{W}vAlign").get(f"{W}val") == "center"
    spacing = tc.find(f"{W}p/{W}pPr/{W}spacing")
    assert spacing.get(f"{W}after") == "60"  # 3 pt


# ── the fonts LibreOffice is given ──────────────────────────────────────


def test_the_conversion_profile_holds_the_bundled_faces(tmp_path):
    copied = give_soffice_the_bundled_fonts(tmp_path)
    assert {"IBMPlexSansArabic-Regular.ttf", "IBMPlexSansArabic-Bold.ttf"} <= set(copied)
    assert (tmp_path / "user" / "fonts" / "IBMPlexSansArabic-Regular.ttf").is_file()


def test_the_pdf_route_converts_with_that_profile(tmp_path, monkeypatch):
    """The DOCX→LibreOffice route hands soffice a profile holding the bundle."""
    from kazma_core.documents import binaries
    from kazma_core.documents.engines import pdf as pdfmod

    seen: dict[str, list[str]] = {}

    def fake_soffice(args, **kwargs):
        profile = next(a for a in args if a.startswith("-env:UserInstallation="))
        from urllib.parse import unquote, urlparse

        path = Path(unquote(urlparse(profile.split("=", 1)[1]).path).lstrip("/") if os.name == "nt"
                    else unquote(urlparse(profile.split("=", 1)[1]).path))
        seen["fonts"] = sorted(f.name for f in (path / "user" / "fonts").glob("*.ttf"))
        raise RuntimeError("stop here")

    monkeypatch.setattr(binaries, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(binaries, "run_soffice_cli", fake_soffice)
    model = ContentModel()
    model.add(BodyBlock(text="نص عربي"))
    pdfmod.PdfEngine(DocProfile.for_content("نص عربي"), []).render(model, tmp_path / "x.pdf")
    assert "IBMPlexSansArabic-Regular.ttf" in seen["fonts"]


# ── the real worker image (opt-in) ──────────────────────────────────────

IMAGE = os.environ.get("KAZMA_TEST_DOCUMENT_WORKER_IMAGE", "")


@pytest.mark.skipif(not IMAGE or not shutil.which("docker"), reason="needs KAZMA_TEST_DOCUMENT_WORKER_IMAGE and docker")
def test_libreoffice_in_the_worker_image_renders_in_plex_only_with_the_bundle(tmp_path):
    """The live failure, in the image: no bundle in the profile -> DejaVu and
    Noto Serif (the negative control); the bundle -> Plex."""
    import pymupdf

    model = ContentModel()
    model.add(TitleBlock(text="أرشيف تغريدات كاظمه على X", level=0))
    model.add(BodyBlock(text=ARCHIVE))
    DocxEngine(DocProfile.for_content(ARCHIVE)).render(model, tmp_path / "a.docx")
    give_soffice_the_bundled_fonts(tmp_path / "with")
    (tmp_path / "without").mkdir()
    fonts = {}
    for profile in ("with", "without"):
        subprocess.run(["docker", "run", "--rm", "--network", "none", "-v", f"{tmp_path}:/w",
                        "--entrypoint", "soffice", IMAGE, "--headless", "--nologo", "--norestore",
                        f"-env:UserInstallation=file:///w/{profile}", "--convert-to", "pdf",
                        "--outdir", f"/w/{profile}", "/w/a.docx"], check=True, timeout=300,
                       capture_output=True)
        doc = pymupdf.open(tmp_path / profile / "a.pdf")
        fonts[profile] = {f[3].split("+")[-1] for page in doc for f in page.get_fonts()}
    assert fonts["with"] == {"IBMPlexSansArabic", "IBMPlexSansArabic-Bold"}, fonts
    assert not any("Plex" in f for f in fonts["without"]), fonts


# ── left-to-right runs inside right-to-left text ────────────────────────

LRM = "\u200e"
LRI, PDI = "\u2066", "\u2069"
DELETED = "10 · 2026-10-05 — 15:15 UTC · 2107127443748266448 — محذوفة"


def _shown(text: str) -> str:
    return text.replace(LRM, "|")


@pytest.mark.parametrize(("text", "isolated"), [
    (DELETED, "|10 · 2026-10-05 — 15:15 UTC · 2107127443748266448| — محذوفة"),
    # a number alone keeps the Arabic convention (W2), which the readers expect
    ("تاريخ التوثيق: 2026-10-18", "تاريخ التوثيق: 2026-10-18"),
    ("الحساب: @KazmaAI · المصدر", "الحساب: |@KazmaAI| · المصدر"),
    ("بتوقيت الكويت (+3).", "بتوقيت الكويت (+3)."),
    ("الساعة +3 UTC هنا", "الساعة |+3 UTC| هنا"),
    ("**قيد:** أحدث N سجلاً", "**قيد:** أحدث |N| سجلاً"),
    ("كلمة-كلمة #كاظمه", "كلمة-كلمة #كاظمه"),
    ("Latin only 2026-10-18", "Latin only 2026-10-18"),
])
def test_each_ltr_run_in_rtl_text_is_isolated(text, isolated):
    from kazma_core.documents.arabic import isolate_ltr_runs

    assert _shown(isolate_ltr_runs(text)) == isolated


def test_the_docx_isolates_ltr_runs_in_rtl_blocks_only(tmp_path):
    root = _docx(tmp_path, f"## {DELETED}\n\nتاريخ التوثيق: 2026-10-18\n\nKazma, counted 2026-10-16")
    texts = [t for _, t in _paragraphs(root)]
    assert f"{LRM}10 · 2026-10-05 — 15:15 UTC · 2107127443748266448{LRM} — محذوفة" in texts
    assert "تاريخ التوثيق: 2026-10-18" in texts  # a number alone is not isolated
    assert "Kazma, counted 2026-10-16" in texts  # an LTR block is left as written


def test_each_mark_is_drawn_in_a_face_that_has_it(tmp_path):
    """Plex has no mark glyph: a mark left in a Plex run is drawn with the
    space glyph, and the PDF's text map names that face's spaces by it."""
    from kazma_core.documents.fonts import direction_mark_family

    family = direction_mark_family()
    assert family, "the bundle carries a face with the direction marks (Amiri)"
    root = _docx(tmp_path, f"## {DELETED}")
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    marked = [r for r in root.iter(f"{w}r") if LRM in "".join(t.text or "" for t in r.iter(f"{w}t"))]
    assert marked and len(marked) % 2 == 0, "each run is marked on both sides"
    for run in marked:
        assert "".join(t.text or "" for t in run.iter(f"{w}t")) == LRM, "a mark is a run of its own"
        fonts = run.find(f"{w}rPr/{w}rFonts")
        assert {fonts.get(f"{w}{slot}") for slot in ("ascii", "hAnsi", "cs")} == {family}


def test_kazma_reads_its_own_pdf_without_the_marks():
    from kazma_core.documents.arabic import isolate_ltr_runs, to_logical

    assert to_logical(isolate_ltr_runs(DELETED)) == DELETED


def _word_x(pdf: Path, words: list[str]) -> list[float]:
    import pymupdf

    doc = pymupdf.open(pdf)
    page_words = [w for page in doc for w in page.get_text("words")]
    # The words on the line of the last one asked for (the heading's id).
    line_y = next(w[1] for w in page_words if w[4] == words[-1])
    found = {w[4]: w[0] for w in page_words if abs(w[1] - line_y) < 2}
    return [found[w] for w in words]


def _old_isolates(self, text: str) -> str:
    """The isolation of 2026-10-18 morning: LRI ... PDI, left in Plex runs."""
    from kazma_core.documents import arabic

    marked = arabic.isolate_ltr_runs(text) if self.profile.block_is_rtl(text) else text
    return marked.replace(LRM, LRI, 1).replace(LRM, PDI, 1)


def _soffice_pdf(tmp_path: Path, monkeypatch, *, isolate: bool | str) -> Path:
    from kazma_core.documents.binaries import find_soffice, run_soffice_cli

    if not isolate:
        monkeypatch.setattr(DocxEngine, "_isolated", lambda self, text: text)
    if isolate == "old":
        monkeypatch.setattr(DocxEngine, "_isolated", _old_isolates)
        monkeypatch.setattr(DocxEngine, "_give_direction_marks_a_face", lambda self, document: None)
    out = tmp_path / {True: "on", False: "off", "old": "old"}[isolate]
    out.mkdir()
    model = ContentModel()
    model.add(BodyBlock(text=(
        f"## {DELETED}\n\nتاريخ التوثيق: 2026-10-18\n\n"
        "| البند | القيمة |\n|---|---|\n| السجلات المقروءة | 3 |\n| حد الأحرف | 280 |\n\nبعد الجدول"
    )))
    DocxEngine(DocProfile.for_content(DELETED)).render(model, out / "a.docx")
    give_soffice_the_bundled_fonts(out / "profile")
    assert find_soffice()
    run_soffice_cli(("--headless", "--nologo", "--norestore",
                     f"-env:UserInstallation={(out / 'profile').as_uri()}",
                     "--convert-to", "pdf", "--outdir", str(out), str(out / "a.docx")),
                    timeout=120, cwd=out)
    return out / "a.pdf"


def _soffice_available() -> bool:
    try:
        import pymupdf  # noqa: F401
        from kazma_core.documents.binaries import find_soffice

        return bool(find_soffice())
    except Exception:
        return False


@pytest.mark.skipif(not _soffice_available(), reason="needs LibreOffice and PyMuPDF")
def test_libreoffice_draws_the_heading_and_the_date_in_reading_order(tmp_path, monkeypatch):
    """On the page: "10" left of the date left of the id (one LTR unit);
    without the isolates (the negative control) LibreOffice drew the groups
    reversed."""
    order = ["10", "2026-10-05", "2107127443748266448"]
    pdf = _soffice_pdf(tmp_path, monkeypatch, isolate=True)
    xs = _word_x(pdf, order)
    assert xs == sorted(xs), xs

    old = _soffice_pdf(tmp_path, monkeypatch, isolate=False)
    old_xs = _word_x(old, order)
    assert old_xs != sorted(old_xs), "the old drawing reversed the groups"


@pytest.mark.skipif(not _soffice_available(), reason="needs LibreOffice and PyMuPDF")
def test_the_pdfs_text_reads_back_with_its_spaces(tmp_path, monkeypatch):
    """Live 2026-10-18: the archive's bold headings read back with U+2066 for
    every space (the isolates drawn with Plex's space glyph), and Kazma's own
    reader failed on the page and fell back to a garbled reading."""
    from kazma_core.documents.service import DocumentService

    marks = "\u200e\u200f\u2066\u2067\u2068\u2069"
    pdf = _soffice_pdf(tmp_path, monkeypatch, isolate=True)
    import pymupdf

    layer = "".join(page.get_text() for page in pymupdf.open(pdf))
    assert "محذوفة" in layer and "UTC" in layer
    assert not any(c in layer.replace("\u200e", "") for c in marks), "no space is named by a mark"
    ir = DocumentService().parse_ingested_blob(pdf, force_ocr=False, ocr_language="ara")
    read = " ".join(b.text for page in ir.pages for b in page.blocks)
    assert "15:15 UTC" in read and "تاريخ التوثيق" in read, read
    assert read.count("حد الأحرف") == 1, "the table is read once"
    assert read.index("حد الأحرف") < read.index("بعد الجدول"), "where it stands"

    old = _soffice_pdf(tmp_path, monkeypatch, isolate="old")
    old_layer = "".join(page.get_text() for page in pymupdf.open(old))
    assert LRI in old_layer or PDI in old_layer, "the old marks named the spaces"


# ── reading a document back: a bidi mark is not a cut ──────────────────


def test_a_paragraph_with_a_bidi_mark_does_not_end_the_parse(tmp_path):
    """``to_logical`` strips bidi controls, so the block came back shorter
    and ``IRBuilder.add_page`` took it for a limit cut: every paragraph after
    the first one holding an RLM or an isolate was dropped, with no warning."""
    import docx
    from kazma_core.documents.service import DocumentService

    source = tmp_path / "marks.docx"
    document = docx.Document()
    document.add_paragraph("مرحبا \u200f بكم في التقرير")
    document.add_paragraph("second paragraph after the mark")
    document.add_paragraph("third paragraph")
    document.save(source)
    ir = DocumentService().parse_ingested_blob(source, force_ocr=False, ocr_language="ara")
    texts = [b.text for page in ir.pages for b in page.blocks]
    assert texts[-1] == "third paragraph", texts
    assert "\u200f" not in texts[0]
