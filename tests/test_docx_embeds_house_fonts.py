"""A Word file Kazma makes carries the house fonts.

Live 2026-10-22: the owner converted an OCR'd booklet to Word and the file
"did not adopt the Kazma house style". It had: every run named IBM Plex Sans
Arabic -- but a DOCX names its fonts, it does not carry them, so on any device
without Plex installed Word drew another face. The PDF route embeds its fonts
at render time; the DOCX engine now embeds the bundled faces as Word does
(obfuscated font parts, ECMA-376 17.8.1).
"""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pytest
from kazma_core.documents.content_model import BodyBlock, ContentModel
from kazma_core.documents.engines.docx import DocxEngine
from kazma_core.documents.engines.docx_fonts import deobfuscate
from kazma_core.documents.fonts import _bundled_pair
from kazma_core.documents.profile import DocProfile

ARABIC = "المغذيات هي مواد يحتاجها الجسم للنمو والطاقة."
FIXTURE = Path(__file__).parent / "fixtures" / "word_embedded_font_sample.json"


def _render(tmp_path: Path, *, embed: bool = True) -> Path:
    model = ContentModel()
    model.add(BodyBlock(text=ARABIC))
    model.add(BodyBlock(text="**English** line, in the same face."))
    engine = DocxEngine(DocProfile.for_content(ARABIC))
    if not embed:
        engine._embed_house_fonts = lambda output: None  # the old engine
    out = tmp_path / "a.docx"
    engine.render(model, out)
    return out


def _embedded(docx: Path) -> dict[str, bytes]:
    """element (embedRegular/embedBold) of IBM Plex Sans Arabic -> font bytes."""

    with zipfile.ZipFile(docx) as z:
        table = z.read("word/fontTable.xml").decode("utf-8")
        rels = z.read("word/_rels/fontTable.xml.rels").decode("utf-8") \
            if "word/_rels/fontTable.xml.rels" in z.namelist() else ""
        entry = re.search(r'<w:font w:name="IBM Plex Sans Arabic">(.*?)</w:font>', table, re.S)
        found = {}
        for element, rel_id, key in re.findall(
                r'<w:(embed\w+) r:id="([^"]+)" w:fontKey="([^"]+)"/>', entry.group(1) if entry else ""):
            target = re.search(rf'Id="{rel_id}"[^>]*Target="([^"]+)"', rels).group(1)
            found[element] = deobfuscate(z.read("word/" + target), key)
        return found


def test_a_generated_word_file_carries_the_bundled_house_faces(tmp_path):
    pair = _bundled_pair(arabic=True)
    if pair is None:
        pytest.skip("the bundled house fonts are not installed with this build")
    regular, bold = pair
    docx = _render(tmp_path)
    found = _embedded(docx)
    assert found.get("embedRegular") == regular.read_bytes()
    assert found.get("embedBold") == bold.read_bytes()
    with zipfile.ZipFile(docx) as z:
        assert "<w:embedTrueTypeFonts/>" in z.read("word/settings.xml").decode("utf-8")
        assert 'Extension="odttf"' in z.read("[Content_Types].xml").decode("utf-8")
        # Still a document python-docx (and so Word) opens.
        assert z.testzip() is None
    import docx as python_docx

    assert ARABIC in "\n".join(p.text for p in python_docx.Document(str(docx)).paragraphs)


def test_negative_control_the_old_engine_names_the_face_and_carries_nothing(tmp_path):
    docx = _render(tmp_path, embed=False)
    with zipfile.ZipFile(docx) as z:
        assert not [n for n in z.namelist() if n.endswith(".odttf")]
    assert _embedded(docx) == {}


def test_the_encoding_is_words_own():
    """Word 16 embedded the same font; its stored head decodes with our key rule."""

    sample = json.loads(FIXTURE.read_text(encoding="utf-8"))
    pair = _bundled_pair(arabic=True)
    if pair is None:
        pytest.skip("the bundled house fonts are not installed with this build")
    font = pair[0].read_bytes()
    stored = bytes.fromhex(sample["obfuscated_head_hex"])
    assert deobfuscate(stored + font[len(stored):], sample["fontKey"]) == font
    assert deobfuscate(font, sample["fontKey"])[: len(stored)] == stored  # its own inverse
    # Negative control: the key's bytes in their written order do not decode it.
    wrong = sample["fontKey"].strip("{}").replace("-", "")
    swapped = "{" + "".join(reversed([wrong[i:i + 2] for i in range(0, 32, 2)])) + "}"
    assert deobfuscate(stored + font[len(stored):], swapped) != font


def test_a_failed_embed_leaves_the_file_whole(tmp_path, monkeypatch):
    import kazma_core.documents.engines.docx_fonts as docx_fonts

    def boom(*args, **kwargs):
        raise OSError("disk")

    monkeypatch.setattr(docx_fonts, "embed_fonts", boom)
    docx = _render(tmp_path)
    import docx as python_docx

    assert ARABIC in "\n".join(p.text for p in python_docx.Document(str(docx)).paragraphs)
