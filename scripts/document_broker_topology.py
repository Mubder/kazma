"""Qualify the document worker broker's Docker topology end to end.

Mirrors deploy/docker-compose.document-broker.yml with plain `docker run`:
the broker container alone mounts the engine socket; a client container
standing in for the Kazma app (uid 10001, no network, read-only root, no
engine socket) reaches it through the shared unix-socket volume and runs
DocumentService through the container profile -- inspection, a bilingual
parse, the hostile corpus, a refused wrong token -- then the broker is
stopped with SIGTERM and no job container may remain. Every container and
the volume it creates are removed, pass or fail.

Usage (CI runs this after building both images):
  python scripts/document_broker_topology.py <worker image> <broker image> [receipt.json]

Needs only the docker CLI; Kazma runs inside the containers. Exit 1 when a
check fails, with each check printed.
"""

from __future__ import annotations

import json
import secrets
import subprocess
import sys
import tempfile
import time
import uuid
from pathlib import Path

WORKER, BROKER = sys.argv[1], sys.argv[2]
RECEIPT = Path(sys.argv[3]) if len(sys.argv) > 3 else None
TAG = uuid.uuid4().hex[:8]
NAME = f"f02topo{TAG}"
VOLUME = f"kazma-f02-sock-{TAG}"
BROKER_CONTAINER = f"kazma-f02-broker-{TAG}"

CLIENT = r'''
import json, os
from dataclasses import replace
from pathlib import Path
from docx import Document
from pypdf import PdfWriter
from kazma_core.documents import worker_container
from kazma_core.documents.config import DocumentConfig
from kazma_core.documents.errors import DocumentParseError
from kazma_core.documents.hostile_corpus import write_hostile_corpus
from kazma_core.documents.service import DocumentService

root = Path("/tmp/client"); root.mkdir()
pdf = root / "blank.pdf"; w = PdfWriter(); w.add_blank_page(width=100, height=100); w.write(str(pdf))
docx = root / "bilingual.docx"; d = Document(); d.add_paragraph("English line 1,234."); d.add_paragraph("سطر عربي ٥٦٧."); d.save(str(docx))
cfg = replace(DocumentConfig(storage_root=root / "store"), ocr_enabled=False, security_local_salvage=False)
svc = DocumentService(config=cfg)
out = {
    "client_uid": os.getuid(),
    "client_engine_socket_visible": any(os.path.exists(p) for p in ("/var/run/docker.sock", "/run/docker.sock")),
    "isolation_report": worker_container.isolation_report(),
    "inspect_pdf": svc.inspect_path(pdf).mime_type,
    "docx_texts": [b.text for p in svc._parse_isolated(docx).pages for b in p.blocks],
}
_, manifest = write_hostile_corpus(root / "corpus")
outcomes = {}
for case in manifest["cases"]:
    if case["stage"] == "sandbox":
        continue
    service = DocumentService(config=replace(cfg, storage_root=root / case["id"], **case["config_overrides"]))
    path = root / "corpus" / case["filename"]
    try:
        service.inspect_path(path) if case["stage"] == "sniff" else service._parse_isolated(path)
        code = "ok"
    except DocumentParseError as exc:
        code = exc.code
    outcomes[case["id"]] = {"code": code, "expected": case["expected_codes"] or ["ok"]}
out["hostile_corpus"] = outcomes
os.environ["KAZMA_DOCUMENT_BROKER_TOKEN_FILE"] = "/tmp/wrong-token"
Path("/tmp/wrong-token").write_text("w" * 48)
try:
    svc.inspect_path(pdf)
    out["wrong_token"] = "accepted"
except DocumentParseError as exc:
    out["wrong_token"] = exc.code
print("KZRESULT" + json.dumps(out, ensure_ascii=False))
'''


def docker(*args: str, check: bool = True, timeout: float = 300) -> subprocess.CompletedProcess[str]:
    return subprocess.run(["docker", *args], capture_output=True, text=True, encoding="utf-8",
                          errors="replace", check=check, timeout=timeout)


def labelled() -> list[str]:
    return docker("ps", "-a", "-q", "--filter", f"label=io.kazma.document-worker={NAME}").stdout.split()


receipt: dict[str, object] = {"worker_image": WORKER, "broker_image": BROKER, "broker_name": NAME}
with tempfile.TemporaryDirectory() as scratch:
    token = Path(scratch) / "document_broker_token"
    token.write_text(secrets.token_urlsafe(48), encoding="utf-8")
    docker("volume", "create", VOLUME)
    try:
        started = time.monotonic()
        docker(
            "run", "-d", "--name", BROKER_CONTAINER,
            "--user", "0:0", "--group-add", "10001", "--cap-drop", "ALL",
            "--security-opt", "no-new-privileges:true", "--network", "none", "--read-only",
            "--tmpfs", "/tmp:rw,nosuid,nodev,noexec,size=16m",
            "-v", "/var/run/docker.sock:/var/run/docker.sock",
            "-v", f"{VOLUME}:/run/kazma-broker",
            "-v", f"{token}:/run/secrets/document_broker_token:ro",
            BROKER, "--image", WORKER, "--token-file", "/run/secrets/document_broker_token",
            "--listen", "unix:/run/kazma-broker/broker.sock", "--socket-gid", "10001", "--name", NAME,
        )
        logs = ""
        while time.monotonic() - started < 120:
            logs = docker("logs", BROKER_CONTAINER, check=False).stderr
            if "[broker] serving" in logs or "refusing" in logs or "cannot" in logs:
                break
            time.sleep(1)
        receipt["broker_start_seconds"] = round(time.monotonic() - started, 2)
        receipt["broker_serving"] = "[broker] serving" in logs
        client = docker(
            "run", "--rm", "--user", "10001:10001", "--network", "none", "--read-only",
            "--cap-drop", "ALL", "--security-opt", "no-new-privileges:true",
            "--tmpfs", "/tmp:rw,nosuid,nodev,size=128m",
            "-v", f"{VOLUME}:/run/kazma-broker",
            "-v", f"{token}:/run/secrets/document_broker_token:ro",
            "-e", "KAZMA_DOCUMENT_WORKER_ISOLATION=container",
            "-e", "KAZMA_DOCUMENT_BROKER_URL=unix:///run/kazma-broker/broker.sock",
            "-e", "KAZMA_DOCUMENT_BROKER_TOKEN_FILE=/run/secrets/document_broker_token",
            "--entrypoint", "python", WORKER, "-I", "-c", CLIENT,
            check=False, timeout=900,
        )
        receipt["client_exit"] = client.returncode
        marker = next((line for line in client.stdout.splitlines() if line.startswith("KZRESULT")), None)
        receipt["client"] = json.loads(marker[len("KZRESULT"):]) if marker else {"stderr": client.stderr[-4000:]}
        receipt["job_containers_before_stop"] = len(labelled())
        stop_started = time.monotonic()
        docker("stop", "-t", "60", BROKER_CONTAINER, timeout=120)
        receipt["broker_stop_seconds"] = round(time.monotonic() - stop_started, 2)
        state = json.loads(docker("inspect", "--format", "{{json .State}}", BROKER_CONTAINER).stdout)
        receipt["broker_exit_code"] = state.get("ExitCode")
        tail = docker("logs", BROKER_CONTAINER, check=False).stderr.splitlines()
        receipt["broker_log_tail"] = [line[-240:] for line in tail[-12:]]
        receipt["job_containers_after_stop"] = len(labelled())
    finally:
        docker("rm", "-f", BROKER_CONTAINER, check=False)
        for cid in labelled():
            docker("rm", "-f", cid, check=False)
        docker("volume", "rm", VOLUME, check=False)
    receipt["leftover_job_containers"] = len(labelled())

client_report = receipt.get("client") or {}
corpus = client_report.get("hostile_corpus") or {}
receipt["checks"] = {
    "broker_serving": receipt.get("broker_serving") is True,
    "client_has_no_engine_socket": client_report.get("client_engine_socket_visible") is False,
    "client_runs_as_10001": client_report.get("client_uid") == 10001,
    "isolation_ready": (client_report.get("isolation_report") or {}).get("ready") is True,
    "pdf_inspected": client_report.get("inspect_pdf") == "application/pdf",
    "arabic_text_parsed": any("سطر عربي" in text for text in client_report.get("docx_texts", [])),
    "hostile_corpus_expected": bool(corpus) and all(v["code"] in v["expected"] for v in corpus.values()),
    "wrong_token_refused": client_report.get("wrong_token") == "worker_isolation_unavailable",
    "no_job_containers_left": receipt.get("leftover_job_containers") == 0,
    "broker_stopped_cleanly": receipt.get("broker_exit_code") == 0,
}
receipt["passed"] = all(receipt["checks"].values())
if RECEIPT is not None:
    RECEIPT.write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(receipt["checks"], indent=2))
sys.exit(0 if receipt["passed"] else 1)
